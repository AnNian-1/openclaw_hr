#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""docx_to_json.py — 将 360 反馈 Word 表解析为标准 JSON

对齐 blank 模板（models/360-feedback-form-blank.docx）填写后的 DOCX 结构，
产出与原 raw/*.json 同契约的反馈 JSON，供 check_feedback / aggregate 使用。

可作为模块 import，也可 CLI：
  python3 docx_to_json.py <docx_dir> [out_dir]
"""
from __future__ import annotations

import json
import os
import re
import sys
from collections import defaultdict
from pathlib import Path


def _reexec_openclaw_hr():
    try:
        import docx  # noqa: F401
        return
    except ImportError:
        pass
    py = os.path.expanduser("~/miniconda3/envs/openclaw-hr/bin/python3")
    if os.path.isfile(py) and os.path.realpath(py) != os.path.realpath(sys.executable):
        os.execv(py, [py] + sys.argv)


_reexec_openclaw_hr()
from docx import Document
from docx.oxml.ns import qn

DIM_KEYS = [
    "goal_achievement",
    "professional_skill",
    "collaboration",
    "responsibility",
    "growth",
]

RELATION_FROM_LABEL = {
    "上级": "manager",
    "平级同事": "peer",
    "平级": "peer",
    "下属": "report",
    "本人自评": "self",
    "自评": "self",
}


def cell_text(cell) -> str:
    parts = []
    for p in cell.paragraphs:
        t = "".join(r.text for r in p.runs).strip()
        if t:
            parts.append(t)
    return "\n".join(parts).strip()


def get_sdt_display(doc, tag: str) -> str:
    for sdt in doc.element.body.iter(qn("w:sdt")):
        tag_el = sdt.find(f".//{qn('w:tag')}")
        if tag_el is None or tag_el.get(qn("w:val")) != tag:
            continue
        content = sdt.find(qn("w:sdtContent"))
        if content is None:
            continue
        texts = content.findall(f".//{qn('w:t')}")
        return "".join((t.text or "") for t in texts).strip()
    raise KeyError(f"SDT tag not found: {tag}")


def parse_period(display: str) -> str:
    s = display.replace("▼", "").strip()
    s = re.sub(r"\s+", "", s)
    m = re.search(r"(20\d{2})\s*[-–—]?\s*[HhＨ]?\s*([1-4])", s, re.I)
    if not m:
        m = re.search(r"(20\d{2})H([1-4])", s, re.I)
    if not m:
        raise ValueError(f"无法解析考核周期: {display!r}")
    return f"{m.group(1)}-H{m.group(2)}"


def parse_relation(display: str) -> str:
    s = display.replace("▼", "").strip()
    s = re.sub(r"^请选择（单选）\s*", "", s).strip()
    for label, key in RELATION_FROM_LABEL.items():
        if label in s or s == label:
            return key
    raise ValueError(f"无法解析关系: {display!r}")


def parse_score(display: str) -> int:
    s = display.replace("▼", "").strip()
    m = re.match(r"^([1-5])", s)
    if not m:
        raise ValueError(f"无法解析评分: {display!r}")
    return int(m.group(1))


def parse_cn_date(text: str) -> str:
    m = re.search(r"(20\d{2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", text)
    if not m:
        raise ValueError(f"无法解析提交日期: {text!r}")
    return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"


def clean_comment(text: str) -> str:
    t = text.strip()
    if t.startswith(("（请写具体事例", "（本周期最值得", "（希望对方")):
        return ""
    return t


def subject_slug(subject: str) -> str:
    """被评估人 → 落盘子目录名（ASCII，供路径与 Canvas id）。

    从 DOCX「被评估人」字段抽取，中文名转拼音拼接（如 张磊→zhanglei），
    不维护姓名白名单。已是拉丁字母的名字则直接清洗小写。
    """
    subject = (subject or "").strip()
    if not subject:
        return "unknown"

    # 已无汉字：按拉丁清洗
    if not re.search(r"[\u4e00-\u9fff]", subject):
        ascii_slug = re.sub(r"[^a-zA-Z0-9]+", "", subject).lower()
        return ascii_slug or "unknown"

    try:
        from pypinyin import Style, lazy_pinyin
    except ImportError as e:
        raise ImportError(
            "解析中文被评估人需要 pypinyin：pip install pypinyin"
        ) from e

    def _keep_latin(ch: str) -> str:
        # 非汉字片段：只留字母数字
        return re.sub(r"[^a-zA-Z0-9]+", "", ch).lower()

    parts = lazy_pinyin(subject, style=Style.NORMAL, errors=_keep_latin)
    slug = "".join(p for p in parts if p).lower()
    slug = re.sub(r"[^a-z0-9]+", "", slug)
    return slug or "unknown"


def subject_dir_name(subject: str) -> str:
    """对外别名，语义同 subject_slug。"""
    return subject_slug(subject)


def parse_docx(path: Path) -> dict:
    """解析单份 DOCX，返回反馈 dict（含临时键 _source_docx）。"""
    doc = Document(str(path))
    if len(doc.tables) < 15:
        raise ValueError(
            f"{path.name}: 表格数量不足（期望 ≥15，实际 {len(doc.tables)}）"
        )

    t0 = doc.tables[0]
    subject = cell_text(t0.rows[0].cells[1])
    subject_role = cell_text(t0.rows[1].cells[1])
    reviewer_alias = cell_text(t0.rows[3].cells[1])

    period = parse_period(get_sdt_display(doc, "period"))
    relation = parse_relation(get_sdt_display(doc, "relation"))

    scores, comments = {}, {}
    for i, key in enumerate(DIM_KEYS):
        scores[key] = parse_score(get_sdt_display(doc, f"score_{key}"))
        comments[key] = clean_comment(
            cell_text(doc.tables[3 + i * 2].rows[0].cells[0])
        )

    strengths = clean_comment(cell_text(doc.tables[12].rows[0].cells[0]))
    improvements = clean_comment(cell_text(doc.tables[13].rows[0].cells[0]))
    submitted_at = parse_cn_date(cell_text(doc.tables[14].rows[0].cells[1]))

    return {
        "subject": subject,
        "subject_role": subject_role,
        "period": period,
        "reviewer_relation": relation,
        "reviewer_alias": reviewer_alias,
        "submitted_at": submitted_at,
        "scores": scores,
        "comments": comments,
        "strengths": strengths,
        "improvements": improvements,
        "_source_docx": path.name,
    }


def assign_json_names(rows: list[dict]) -> None:
    """就地写入 _json_name / feedback_id。

    完全按 DOCX 内「与被评估人关系」下拉值分桶命名（不依赖文件名）：
      上级 → manager.json
      自评 → self.json
      平级 → peer-01.json, peer-02.json, …
      下属 → report-01.json, report-02.json, …
    同关系多份时按评估人别名 + 源文件名排序，保证可复现。
    """
    buckets = defaultdict(list)
    for row in rows:
        rel = row.get("reviewer_relation")
        if rel not in ("manager", "peer", "report", "self"):
            raise ValueError(
                f"{row.get('_source_docx')}: 无法识别评估关系 "
                f"（reviewer_relation={rel!r}），请检查 DOCX「与被评估人关系」"
            )
        buckets[rel].append(row)

    used = set()
    for relation, items in buckets.items():
        items.sort(
            key=lambda r: (
                (r.get("reviewer_alias") or ""),
                (r.get("_source_docx") or ""),
            )
        )
        if relation in ("manager", "self"):
            if len(items) != 1:
                names = ", ".join(r.get("_source_docx", "?") for r in items)
                raise ValueError(
                    f"{relation} 关系应有且仅有 1 份，实际 {len(items)} 份：{names}"
                )
            candidate = f"{relation}.json"
            if candidate in used:
                raise ValueError(f"重复的 {relation} 反馈：已有 {candidate}")
            items[0]["_json_name"] = candidate
            used.add(candidate)
            continue
        for i, item in enumerate(items, start=1):
            candidate = f"{relation}-{i:02d}.json"
            if candidate in used:
                raise ValueError(f"命名冲突：{candidate}")
            item["_json_name"] = candidate
            used.add(candidate)

    for row in rows:
        json_name = row["_json_name"]
        period_slug = row["period"].lower().replace("-", "")
        suffix = Path(json_name).stem.replace("-", "")
        row["feedback_id"] = (
            f"fb-360-{subject_slug(row['subject'])}-{period_slug}-{suffix}"
        )


def convert_dir(docx_dir: Path, out_dir: Path | None = None) -> list[Path]:
    """解析目录下全部 DOCX，写入 out_dir（默认同目录）。返回写出的 JSON 路径。"""
    docx_dir = Path(docx_dir)
    out_dir = Path(out_dir) if out_dir else docx_dir
    out_dir.mkdir(parents=True, exist_ok=True)

    docx_files = sorted(
        p for p in docx_dir.rglob("*.docx")
        if not p.name.startswith("~$") and not p.name.startswith(".")
    )
    if not docx_files:
        raise FileNotFoundError(f"{docx_dir} 下没有 .docx")

    rows = [parse_docx(p) for p in docx_files]
    assign_json_names(rows)

    written = []
    for row in rows:
        json_name = row.pop("_json_name")
        row.pop("_source_docx", None)
        out = out_dir / json_name
        out.write_text(
            json.dumps(row, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        written.append(out)
        print(f"  wrote {json_name}")
    return written


def main(argv=None):
    argv = list(argv or sys.argv[1:])
    if not argv:
        print("用法: python3 docx_to_json.py <docx_dir> [out_dir]", file=sys.stderr)
        sys.exit(2)
    docx_dir = Path(argv[0])
    out_dir = Path(argv[1]) if len(argv) > 1 else docx_dir
    if not docx_dir.is_dir():
        print(f"目录不存在: {docx_dir}", file=sys.stderr)
        sys.exit(2)
    paths = convert_dir(docx_dir, out_dir)
    print(f"done → {out_dir} ({len(paths)} json)")


if __name__ == "__main__":
    main()
