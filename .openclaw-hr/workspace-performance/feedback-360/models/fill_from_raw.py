#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
按 models/360-feedback-form-blank.docx 模板，
将 raw/*.json 填成可阅读的 DOCX，输出到 source_feedback/zhanglei/。
"""
from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
BLANK = ROOT / "models" / "360-feedback-form-blank.docx"
RAW = ROOT / "raw"
OUT = ROOT / "source_feedback" / "zhanglei"

FONT_BODY = "宋体"
FONT_UI = "微软雅黑"
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
MUTED = RGBColor(0x4A, 0x55, 0x68)

DIM_KEYS = [
    "goal_achievement",
    "professional_skill",
    "collaboration",
    "responsibility",
    "growth",
]

RELATION_LABEL = {
    "manager": "上级",
    "peer": "平级同事",
    "report": "下属",
    "self": "本人自评",
}

SCORE_LABEL = {
    5: "5-远超预期",
    4: "4-超出预期",
    3: "3-符合预期",
    2: "2-部分未达",
    1: "1-明显未达",
}

# raw 文件名 → 输出文件名
FILE_MAP = [
    ("manager.json", "360-反馈表-张磊-上级-王建华.docx"),
    ("peer-01.json", "360-反馈表-张磊-平级-陈思远.docx"),
    ("peer-02.json", "360-反馈表-张磊-平级-刘婷.docx"),
    ("peer-03.json", "360-反馈表-张磊-平级-赵宇.docx"),
    ("report-01.json", "360-反馈表-张磊-下属-孙浩.docx"),
    ("report-02.json", "360-反馈表-张磊-下属-周琳.docx"),
    ("self.json", "360-反馈表-张磊-自评.docx"),
]


def set_run_font(run, size=10.5, bold=False, color=None, name=FONT_BODY):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name if name != "黑体" else "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name if name != "黑体" else "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def clear_paragraph(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)


def set_cell_text(cell, text, *, size=10.5, color=NAVY, name=FONT_UI, bold=False):
    """清空单元格文字并写入；保留边框/底色。"""
    paragraphs = cell.paragraphs
    if not paragraphs:
        return
    clear_paragraph(paragraphs[0])
    for extra in paragraphs[1:]:
        clear_paragraph(extra)
    run = paragraphs[0].add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, name=name)


def set_comment_cell(cell, text):
    """评语框：去掉占位提示，写入正文。"""
    paragraphs = cell.paragraphs
    if not paragraphs:
        return
    clear_paragraph(paragraphs[0])
    paragraphs[0].paragraph_format.space_before = Pt(4)
    paragraphs[0].paragraph_format.space_after = Pt(4)
    paragraphs[0].paragraph_format.line_spacing = 1.45
    run = paragraphs[0].add_run(text)
    set_run_font(run, size=10.5, color=NAVY, name=FONT_BODY)
    for extra in paragraphs[1:]:
        clear_paragraph(extra)


def set_sdt_display(doc, tag: str, display: str):
    """设置下拉内容控件当前显示值。"""
    for sdt in doc.element.body.iter(qn("w:sdt")):
        tag_el = sdt.find(f".//{qn('w:tag')}")
        if tag_el is None or tag_el.get(qn("w:val")) != tag:
            continue
        content = sdt.find(qn("w:sdtContent"))
        if content is None:
            continue
        texts = content.findall(f".//{qn('w:t')}")
        if texts:
            texts[0].text = display
            for t in texts[1:]:
                t.text = ""
        else:
            # 兜底：在 sdtContent 内写一个 run
            from docx.oxml import parse_xml
            from docx.oxml.ns import nsdecls

            content.append(
                parse_xml(
                    f'<w:r {nsdecls("w")}>'
                    f'<w:rPr><w:rFonts w:ascii="{FONT_UI}" w:hAnsi="{FONT_UI}"'
                    f' w:eastAsia="{FONT_UI}" w:cs="{FONT_UI}"/>'
                    f'<w:sz w:val="21"/><w:szCs w:val="21"/>'
                    f'<w:color w:val="1B2A4A"/></w:rPr>'
                    f"<w:t>{_xml_escape(display)}</w:t></w:r>"
                )
            )
        return
    raise KeyError(f"SDT tag not found: {tag}")


def _xml_escape(s: str) -> str:
    return (
        str(s)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def normalize_period(period: str) -> str:
    """2026-H1 / 2026 H1 → 2026H1（对齐模板下拉选项）。"""
    return re.sub(r"[\s\-–—]", "", period).upper().replace("Ｈ", "H")


def short_name(alias: str) -> str:
    """王建华（研发组组长…）→ 王建华"""
    m = re.match(r"^([^（(]+)", alias.strip())
    return m.group(1).strip() if m else alias.strip()


def format_cn_date(iso: str) -> str:
    """2026-06-28 → 2026年06月28日"""
    y, m, d = iso.split("-")
    return f"{y}年{m}月{d}日"


def fill_one(data: dict, out_path: Path):
    shutil.copy2(BLANK, out_path)
    doc = Document(str(out_path))

    period = normalize_period(data["period"])
    relation = data["reviewer_relation"]
    relation_label = RELATION_LABEL[relation]
    scores = data["scores"]
    comments = data["comments"]

    # —— 基本信息 ——
    t0 = doc.tables[0]
    set_cell_text(t0.rows[0].cells[1], data["subject"])
    set_cell_text(t0.rows[1].cells[1], data["subject_role"])
    set_cell_text(t0.rows[3].cells[1], data["reviewer_alias"])

    set_sdt_display(doc, "period", period)
    set_sdt_display(doc, "relation", relation_label)

    # —— 五个维度：评分 + 事实与观察 ——
    # tables: 2/3, 4/5, 6/7, 8/9, 10/11
    for i, key in enumerate(DIM_KEYS):
        score = int(scores[key])
        set_sdt_display(doc, f"score_{key}", SCORE_LABEL[score])
        comment_table = doc.tables[3 + i * 2]
        set_comment_cell(comment_table.rows[0].cells[0], comments[key])

    # —— 总体印象 ——
    set_comment_cell(doc.tables[12].rows[0].cells[0], data["strengths"])
    set_comment_cell(doc.tables[13].rows[0].cells[0], data["improvements"])

    # —— 签名 / 日期 ——
    sig = doc.tables[14]
    left, right = sig.rows[0].cells
    # 保留红色 *，重写标签+姓名
    clear_paragraph(left.paragraphs[0])
    star = left.paragraphs[0].add_run("*")
    set_run_font(star, size=10.5, bold=True, color=RGBColor(0xC4, 0x28, 0x1C), name=FONT_UI)
    lab = left.paragraphs[0].add_run("评估人签名：")
    set_run_font(lab, size=10.5, color=MUTED, name=FONT_BODY)
    name = left.paragraphs[0].add_run(short_name(data["reviewer_alias"]))
    set_run_font(name, size=10.5, color=NAVY, name=FONT_UI)

    clear_paragraph(right.paragraphs[0])
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    star = right.paragraphs[0].add_run("*")
    set_run_font(star, size=10.5, bold=True, color=RGBColor(0xC4, 0x28, 0x1C), name=FONT_UI)
    lab = right.paragraphs[0].add_run("提交日期：")
    set_run_font(lab, size=10.5, color=MUTED, name=FONT_BODY)
    dt = right.paragraphs[0].add_run(format_cn_date(data["submitted_at"]))
    set_run_font(dt, size=10.5, color=NAVY, name=FONT_UI)

    doc.save(str(out_path))


def main():
    if not BLANK.exists():
        raise SystemExit(f"blank template missing: {BLANK}")
    OUT.mkdir(parents=True, exist_ok=True)

    for src_name, out_name in FILE_MAP:
        src = RAW / src_name
        data = json.loads(src.read_text(encoding="utf-8"))
        out_path = OUT / out_name
        fill_one(data, out_path)
        print("wrote", out_path.relative_to(ROOT))

    print(f"done → {OUT} ({len(FILE_MAP)} files)")


if __name__ == "__main__":
    main()
